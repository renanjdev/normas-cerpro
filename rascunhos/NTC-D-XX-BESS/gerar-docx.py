#!/usr/bin/env python3
"""Gera o .docx da NTC-D-XX no layout da norma original da CERPRO."""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# --- paleta da norma CERPRO ---
DBLUE = RGBColor(0x1F, 0x38, 0x64)   # H1/H2, capa
MBLUE = RGBColor(0x2E, 0x75, 0xB6)   # H3
RED   = RGBColor(0xEE, 0x00, 0x00)   # acento "SISTEMAS BESS"
GREY  = RGBColor(0x60, 0x60, 0x60)
CODE  = RGBColor(0x8A, 0x2A, 0x00)

FILES = ['NTC-D-XX-corpo-consolidado.md', 'NTC-D-XX-anexos.md']
OUT   = 'NTC-D-XX_BESS_CERPRO_consolidada_R1.docx'

doc = Document()

# Normal: Calibri (tema), 11pt — corpo justificado, no padrão da norma CERPRO
normal = doc.styles['Normal']
normal.font.name='Calibri'; normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.15
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY

def clean(t):
    """remove espaços supérfluos (duplos, antes de pontuação) preservando os marcadores."""
    t=re.sub(r'[ \t]{2,}',' ',t)
    t=re.sub(r'\s+([,;.:!?])',r'\1',t)
    return t.strip()

# Headings iguais à norma: H1 16 / H2 13 (1F3864) ; H3 12 / H4 11 (2E75B6)
for i,(sz,col) in enumerate([(16,DBLUE),(13,DBLUE),(12,MBLUE),(11,MBLUE)], start=1):
    h=doc.styles[f'Heading {i}']
    h.font.name='Calibri'; h.font.size=Pt(sz); h.font.bold=True; h.font.color.rgb=col
    h.paragraph_format.space_before=Pt(12 if i<3 else 8); h.paragraph_format.space_after=Pt(4)
    h.paragraph_format.keep_with_next=True

# Página A4, margens da norma
for s in doc.sections:
    s.page_width=Cm(21.0); s.page_height=Cm(29.7)
    s.top_margin=Cm(3.17); s.bottom_margin=Cm(3.0); s.left_margin=Cm(2.54); s.right_margin=Cm(2.54)

def _hl(name):
    return {'yellow':WD_COLOR_INDEX.YELLOW,'cyan':WD_COLOR_INDEX.TURQUOISE,
            'magenta':WD_COLOR_INDEX.PINK}[name]

def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),hexc); tcPr.append(sh)

def _marks(p, text, bold, ital):
    """Nível 3: marcadores atômicos (código, ~~tachado~~, «param», [DECISÃO], [VERIFICAR]),
    herdando negrito/itálico do contexto."""
    pat=r'(`[^`]+`|~~.+?~~|«[^»]*»|\[DECISÃO[^\]]*\]|\[VERIFICAR[^\]]*\])'
    for part in re.split(pat, text):
        if not part: continue
        r=p.add_run()
        if part.startswith('`') and part.endswith('`'):
            r.text=part[1:-1]; r.font.name='Consolas'; r.font.size=Pt(9.5); r.font.color.rgb=CODE
        elif part.startswith('~~') and part.endswith('~~'):
            r.text=part[2:-2]; r.font.strike=True
        elif part.startswith('«') and part.endswith('»'):
            r.text=part; r.font.color.rgb=RGBColor(0x70,0x50,0x00); r.font.highlight_color=_hl('yellow')
        elif part.startswith('[DECISÃO'):
            r.text=part; r.bold=True; r.font.highlight_color=_hl('cyan')
        elif part.startswith('[VERIFICAR'):
            r.text=part; r.bold=True; r.font.highlight_color=_hl('magenta')
        else:
            r.text=part
        if bold: r.bold=True
        if ital: r.italic=True

def add_inline(p, text, base_bold=False):
    # Parser recursivo de ênfase: casa **negrito** e *itálico* em qualquer ordem de
    # aninhamento; trechos `código` são preservados (não interpreta * dentro deles).
    # As folhas vão para _marks (marcadores atômicos), herdando negrito/itálico.
    def ital_close(t, start):
        # acha o '*' de fechamento do itálico, pulando spans **negrito** e `código`
        k=start
        while k<len(t):
            if t.startswith('**', k):
                e=t.find('**', k+2); k=(e+2) if e!=-1 else k+2; continue
            if t[k]=='`':
                e=t.find('`', k+1); k=(e+1) if e!=-1 else k+1; continue
            if t[k]=='*': return k
            k+=1
        return -1
    def render(t, bold, ital):
        i=0; n=len(t); buf=''
        def flush():
            nonlocal buf
            if buf: _marks(p, buf, bold, ital); buf=''
        while i<n:
            if t[i]=='`':                                   # protege código
                j=t.find('`', i+1)
                if j!=-1: buf+=t[i:j+1]; i=j+1; continue
            if t.startswith('**', i):
                j=t.find('**', i+2)
                if j!=-1 and t[i+2:j].strip():
                    flush(); render(t[i+2:j], True, ital); i=j+2; continue
            if t[i]=='*':
                j=ital_close(t, i+1)
                if j!=-1 and '\n' not in t[i+1:j] and t[i+1:j].strip():
                    flush(); render(t[i+1:j], bold, True); i=j+1; continue
            buf+=t[i]; i+=1
        flush()
    render(text, base_bold, False)

def flush_table(rows):
    rows=[r for r in rows if not re.match(r'^[\s:\-|]+$','|'.join(r))]
    if not rows: return
    ncol=max(len(r) for r in rows)
    t=doc.add_table(rows=0, cols=ncol); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for ri,r in enumerate(rows):
        cells=t.add_row().cells
        for ci in range(ncol):
            txt=r[ci] if ci<len(r) else ''
            cells[ci].text=''; p=cells[ci].paragraphs[0]; p.paragraph_format.space_after=Pt(2)
            add_inline(p, txt, base_bold=(ri==0))
            if ri==0:
                shade(cells[ci],'1F3864')
                for run in p.runs: run.font.color.rgb=RGBColor(0xFF,0xFF,0xFF); run.bold=True
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def ctr(size, text, color=None, bold=False, italic=False):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text); r.font.size=Pt(size); r.bold=bold; r.italic=italic
    if color is not None: r.font.color.rgb=color
    return p

# ---------- CAPA (réplica da norma) ----------
for _ in range(2): doc.add_paragraph()
ctr(32,'CERPRO',DBLUE,bold=True)
ctr(14,'Cooperativa de Eletrificação Rural da Região de Promissão',DBLUE)
doc.add_paragraph()
ctr(14,'NTC-D-XX',DBLUE,bold=True)
ctr(14,'NORMA TÉCNICA E PADRONIZAÇÃO',bold=True)
doc.add_paragraph()
ctr(15,'REQUISITOS E PROCEDIMENTOS PARA CONEXÃO DE',bold=True)
ctr(17,'SISTEMAS BESS',RED,bold=True)
ctr(12,'(Battery Energy Storage Systems)',RED)
ctr(15,'AO SISTEMA DE DISTRIBUIÇÃO DA CERPRO',bold=True)
doc.add_paragraph()
ctr(13,'Versão de trabalho R1 — Consolidada (corpo + anexos)',bold=True)
ctr(11,'Base: R0 da Engenharia CERPRO + NT-BESS-001/2026',GREY)
ctr(11,'Elaborado por Engenharia CERPRO',GREY)
ctr(11,'Aprovado pelo Grupo Técnico de Padronização',GREY)
for _ in range(2): doc.add_paragraph()
p=doc.add_paragraph(); r=p.add_run('Legenda dos marcadores editoriais:'); r.bold=True; r.font.color.rgb=DBLUE
for txt in ['«parâmetro» — valor a ser cravado pela Engenharia da CERPRO',
            '[DECISÃO CERPRO] — ponto que requer deliberação',
            '[VERIFICAR] — citação regulatória a confirmar antes da publicação']:
    pp=doc.add_paragraph(style='List Bullet'); add_inline(pp, txt)
doc.add_page_break()

# O SUMÁRIO (TOC) é inserido no corpo, na posição do "## SUMÁRIO" do markdown
# (após Título e Controle de Revisões), via add_toc() — ver loop do corpo.
def add_toc():
    p=doc.add_paragraph()
    fld=OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), r'TOC \o "1-3" \h \z \u')
    run=OxmlElement('w:r'); t=OxmlElement('w:t'); t.set(qn('xml:space'),'preserve')
    t.text='Atualize este campo no Word (clique direito ▸ Atualizar Campo) para gerar o sumário.'
    run.append(t); fld.append(run); p._p.append(fld)

# ---------- CABEÇALHO (réplica da norma: tabela 3x2) ----------
TITULO='Requisitos e Procedimentos para Conexão de Sistemas BESS (Battery Energy Storage Systems) ao Sistema de Distribuição da CERPRO'
def kv(cell, label, value, vbold=False):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    if label:
        rl=p.add_run(label); rl.bold=True; rl.font.size=Pt(8.5); rl.font.color.rgb=DBLUE
    rv=p.add_run(value); rv.font.size=Pt(8.5); rv.bold=vbold

sec=doc.sections[0]
hdr=sec.header; hdr.is_linked_to_previous=False
ht=hdr.add_table(rows=3, cols=2, width=Cm(16.0)); ht.style='Table Grid'; ht.alignment=WD_TABLE_ALIGNMENT.CENTER
kv(ht.cell(0,0),'Tipo: ','Norma Técnica e Padronização')
kv(ht.cell(0,1),'', 'NTC-D-XX', vbold=True)
kv(ht.cell(1,0),'Área de Aplicação: ','Distribuição de Energia Elétrica')
kv(ht.cell(1,1),'Versão: ','R1/2026 (consolidada)')
c=ht.cell(2,0).merge(ht.cell(2,1)); kv(c,'Título: ',TITULO)

# ---------- RODAPÉ (réplica da norma: tabela 1x4) ----------
f=sec.footer; f.is_linked_to_previous=False
ft=f.add_table(rows=1, cols=4, width=Cm(16.0)); ft.style='Table Grid'; ft.alignment=WD_TABLE_ALIGNMENT.CENTER
kv(ft.cell(0,0),'Elaborado por: ','CERPRO')
kv(ft.cell(0,1),'Aprovado por: ','Grupo Técnico de Padronização')
kv(ft.cell(0,2),'Data de vigência: ','__/__/2026')
pc=ft.cell(0,3).paragraphs[0]; pc.paragraph_format.space_after=Pt(0)
rr=pc.add_run('Página '); rr.bold=True; rr.font.size=Pt(8.5); rr.font.color.rgb=DBLUE
fp=OxmlElement('w:fldSimple'); fp.set(qn('w:instr'),'PAGE'); pc._p.append(fp)
r2=pc.add_run(' de '); r2.font.size=Pt(8.5)
fn2=OxmlElement('w:fldSimple'); fn2.set(qn('w:instr'),'NUMPAGES'); pc._p.append(fn2)

# ---------- corpo ----------
def add_code_block(buf):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(0.3)
    p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(6)
    pPr=p._p.get_or_add_pPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),'F2F4F7')
    # w:shd deve preceder tabs/spacing/ind/jc/rPr na ordem do schema
    pPr.insert_element_before(sh,'w:tabs','w:suppressAutoHyphens','w:kinsoku','w:wordWrap',
        'w:overflowPunct','w:topLinePunct','w:autoSpaceDE','w:autoSpaceDN','w:bidi',
        'w:adjustRightInd','w:snapToGrid','w:spacing','w:ind','w:contextualSpacing',
        'w:mirrorIndents','w:suppressOverlap','w:jc','w:textDirection','w:textAlignment',
        'w:textboxTightWrap','w:outlineLvl','w:divId','w:cnfStyle','w:rPr','w:sectPr','w:pPrChange')
    r=p.add_run('\n'.join(buf)); r.font.name='Consolas'; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(0x22,0x33,0x44)

def coalesce(text):
    """Junta linhas de continuação (soft-wrap do Markdown) num único parágrafo lógico,
    respeitando blocos de código. Evita que cada linha física vire um parágrafo solto."""
    out=[]; incode=False
    for line in text.split('\n'):
        st=line.strip()
        if st.startswith('```'):
            incode=not incode; out.append(line); continue
        if incode: out.append(line); continue
        if st=='':
            out.append(''); continue
        is_block=bool(re.match(r'(#{1,6}\s|[-*]\s|\d+\.\s|>\s|\||@@IMG:|---$)', st))
        prev=(out[-1] if out else '').strip()
        # só funde em blocos que admitem continuação: parágrafo, lista, citação.
        # NUNCA em título (#), imagem, regra (---), tabela (|) ou fence.
        prev_mergeable=bool(prev) and not re.match(r'(#{1,6}\s|@@IMG:|---$|\||```)', prev)
        if (not is_block) and prev_mergeable:
            out[-1]=out[-1].rstrip()+' '+st          # continuação → anexa ao bloco anterior
        else:
            out.append(line)
    return out

for fi,fn in enumerate(FILES):
    if fi>0: doc.add_page_break()
    tbl=[]; code=None; skip_summary=False
    for line in coalesce(open(fn).read()):
        s=line.rstrip()
        if s.strip().startswith('```'):
            if code is None: code=[]
            else: add_code_block(code); code=None
            continue
        if code is not None:
            code.append(line); continue
        mimg=re.match(r'^@@IMG:(.+)@@$', s.strip())
        if mimg:
            from docx.shared import Inches as _In
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(mimg.group(1), width=_In(5.0)); continue
        if s.strip().startswith('|') and s.strip().endswith('|'):
            tbl.append([c.strip() for c in s.strip().strip('|').split('|')]); continue
        elif tbl: flush_table(tbl); tbl=[]
        if not s.strip(): continue
        m=re.match(r'^(#{1,6})\s+(.*)', s)
        if m:
            skip_summary=False            # qualquer título encerra o bloco do sumário
            htext=clean(m.group(2))
            h=doc.add_heading('', level=min(len(m.group(1)),4))
            add_inline(h, htext, base_bold=True)
            if htext.strip().upper()=='SUMÁRIO':
                add_toc()                 # TOC real no lugar do sumário do markdown
                skip_summary=True         # pula a lista corrida que vem a seguir
                doc.add_page_break()
            continue
        if s.strip()=='---': skip_summary=False; continue   # fim do bloco do sumário
        if skip_summary: continue          # descarta a lista corrida do sumário (substituída pelo TOC)
        if s.lstrip().startswith('> '):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(0.6)
            p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(3)
            pPr=p._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); left=OxmlElement('w:left')
            left.set(qn('w:val'),'single'); left.set(qn('w:sz'),'18'); left.set(qn('w:space'),'8'); left.set(qn('w:color'),'2E75B6')
            pbdr.append(left)
            # w:pBdr deve preceder shd/tabs/spacing/ind/jc/rPr na ordem do schema
            pPr.insert_element_before(pbdr,'w:shd','w:tabs','w:suppressAutoHyphens','w:kinsoku',
                'w:wordWrap','w:overflowPunct','w:topLinePunct','w:autoSpaceDE','w:autoSpaceDN',
                'w:bidi','w:adjustRightInd','w:snapToGrid','w:spacing','w:ind','w:contextualSpacing',
                'w:mirrorIndents','w:suppressOverlap','w:jc','w:textDirection','w:textAlignment',
                'w:textboxTightWrap','w:outlineLvl','w:divId','w:cnfStyle','w:rPr','w:sectPr','w:pPrChange')
            add_inline(p, clean(s.lstrip()[2:]))
            for run in p.runs:
                if run.font.color.rgb is None: run.font.color.rgb=GREY
            continue
        mb=re.match(r'^(\s*)[-*]\s+(.*)', s)
        if mb:
            style='List Bullet' if len(mb.group(1))<2 else 'List Bullet 2'
            try: p=doc.add_paragraph(style=style)
            except KeyError: p=doc.add_paragraph(style='List Bullet')
            p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, clean(mb.group(2))); continue
        mn=re.match(r'^(\s*)\d+\.\s+(.*)', s)
        if mn:
            p=doc.add_paragraph(style='List Number'); p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, clean(mn.group(2))); continue
        # parágrafo de corpo: justificado com recuo de 1ª linha (padrão norma CERPRO)
        p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Cm(1.25)
        add_inline(p, clean(s))
    if tbl: flush_table(tbl)

# ---------- COMENTÁRIOS DE REVISÃO (painel de Revisão do Word) ----------
# (âncora_substring, texto_do_comentário, autor, iniciais)
REVIEW = [
 ('1.1 ESS, SAE e BESS',
  'TERMINOLOGIA: corrigida a inversão SAE↔BESS do R0 (SAE = ESS família; BESS = baterias). '
  'Confirmar que a equipe adota esta hierarquia em todos os documentos da CERPRO.', 'Revisão Normas', 'RN'),
 ('REN 956/2021',
  'VERIFICAR: confirmar o número/objeto desta REN (no R0 aparece como Procedimentos de Distribuição). '
  'Checar se não é a REN 1.000/2021 + PRODIST. Ajustar antes de publicar.', 'Regulação', 'REG'),
 ('6.1 Mini-glossário',
  'Glossário novo p/ separar Zero-Grid (estado) × Anti-exportação (função) × LPI × SCRPI × Hard Limit. '
  'Engenharia: validar se os meios de comprovação descritos batem com a prática.', 'Revisão Normas', 'RN'),
 ('é referência mínima ao tempo',
  'AFASTAMENTO: o valor «3,0 m» (NFPA 855) é placeholder de referência. Engenharia/Bombeiros: '
  'cravar o afastamento adotado e a regra de redução por ensaio UL 9540A.', 'Segurança', 'SEG'),
 ('AVCB (Auto de Vistoria do Corpo de Bombeiros) ou equivalente',
  'DECISÃO CERPRO (definida): AVCB (ou equivalente) e seguro tornam-se DOCUMENTOS OBRIGATÓRIOS, '
  'com cobertura de incêndio, explosão, danos ambientais, RC e danos à rede. Jurídico: detalhar '
  'limites/vigência mínima da apólice no modelo contratual.', 'Jurídico', 'JUR'),
 ('Tabela do item 9.1.1',
  'PARÂMETROS BT: preencher os ajustes (27/59/81 e tempos) conforme Port. INMETRO 515/2023 vigente. '
  'Bloqueio para publicação.', 'Engenharia', 'ENG'),
 ('9.2.6 Anti-ilhamento em MT',
  'CONFLITO: a NTC-D-09 (Tab.3/4/6/7) EXIGE a função 78 (salto de vetor, Ativo 2,0s). A decisão '
  'preliminar de vedar 78 diverge da norma vigente da CERPRO. Proteção/Diretoria: decidir alinhar '
  '(manter 78) ou divergir (vedar e revisar a NTC-D-09).', 'Proteção', 'PROT'),
 ('Tabela 5 — Suportabilidade a desvios de tensão',
  'FRT: valores agora preenchidos a partir da NTC-D-09 (Tab.6/7 - Port. 515/2023). Engenharia: '
  'validar coerência BESS x MMGD e necessidade de curva LVRT específica.', 'Engenharia', 'ENG'),
 ('10.3 Telemetria e Supervisão',
  'PARÂMETROS: confirmar limiares (telemetria «300 kW»; IEC 61850 >500 kW; IEC 62443 >500 kW). '
  'Operação validar integração ao COS.', 'Operação', 'OPER'),
 ('13. GOVERNANÇA DA ANÁLISE',
  'DECISÃO CERPRO: definir a matriz Classe (A–D) × estudos exigíveis e a composição/alçada do '
  'Comitê Interno BESS.', 'Regulação', 'REG'),
 ('14. DISPOSIÇÕES FINAIS E VIGÊNCIA',
  'PARÂMETRO: definir prazo de vigência e regime de transição (ex.: 120 dias). '
  'Atualizar o Controle de Revisões ao publicar.', 'Revisão Normas', 'RN'),
 ('ANEXO A — Arranjos de Conexão',
  'UNIFILARES: versões esquemáticas (lógicas). Engenharia: produzir os unifilares COTADOS '
  '(TC/TP e relações, bitolas, tensões/correntes por trecho) para anexar ao projeto.', 'Engenharia', 'ENG'),
 ('ANEXO F — Acordo Operativo',
  'JURÍDICO: revisar o modelo de Acordo Operativo (limites operacionais, intervenção da CERPRO, '
  'retenção de registros 60 meses, vínculo com aditivo do CUSD).', 'Jurídico', 'JUR'),
 ('ANEXO D — Lista de Certificações',
  'Engenharia: validar a lista de certificações por componente e definir exigências de tradução '
  'técnica p/ documentos estrangeiros.', 'Engenharia', 'ENG'),
]
def find_para(anchor):
    for p in doc.paragraphs:
        if anchor in p.text and p.runs:
            return p
    return None
added=0
for anchor,text,author,initials in REVIEW:
    p=find_para(anchor)
    if p is None:
        print('  [aviso] âncora não encontrada:', anchor[:40]); continue
    try:
        doc.add_comment(p.runs, text=text, author=author, initials=initials); added+=1
    except Exception as e:
        print('  [erro] comentário em', anchor[:30], '->', e)
print(f'comentários de revisão adicionados: {added}/{len(REVIEW)}')

# corrige w:zoom sem atributo percent (schema OOXML exige percent) — injetado pelo settings padrão
_st=doc.settings.element
for z in _st.findall(qn('w:zoom')):
    if z.get(qn('w:percent')) is None: z.set(qn('w:percent'),'100')

doc.save(OUT)
print('saved', OUT)
